"""Local text extraction for uploaded resumes.

Deliberately dumb and deliberately local: pull the text layer out of a PDF or
docx with no ML involved. A normal digital resume extracts cleanly this way
for free; a scanned resume with no text layer comes back empty, which the
caller treats as "reject, ask for a text-based file" rather than reaching for
OCR or a vision model — see the discussion on issue #4.
"""

from pathlib import Path


class ExtractError(RuntimeError):
    pass


def _extract_pdf(path: Path) -> str:
    import pdfplumber

    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(text)
    return "\n".join(chunks)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Resumes often put skills/experience in tables rather than paragraphs.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def extract_text(path: Path) -> str:
    """Extract whatever text layer the file has. Returns "" for a scanned
    document with no text layer — that is a valid, expected outcome, not an
    error; the caller decides what to do with an empty result."""
    suffix = path.suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        raise ExtractError(f"unsupported file type: {suffix or '(none)'}")
    try:
        return extractor(path).strip()
    except Exception as exc:  # noqa: BLE001 — any parser failure means "couldn't read this file"
        raise ExtractError(f"could not read {suffix} file: {exc}") from exc
